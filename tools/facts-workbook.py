#!/usr/bin/env python3
"""
Builds the companion workbook for the-facts edition on Andover's results.

    python tools/facts-workbook.py the-facts/editions/<edition-id>

Writes andover-edsight-data.xlsx: a documented cover sheet, a key-figures sheet
listing every number cited in the edition with the sheet it can be checked
against, and the underlying data district by district and year by year.

Everything is derived live from the committed EdSight extracts under
data/edsight/, so the workbook cannot drift from the charts or the copy.
Regenerate after any refresh of those datasets.
"""
import csv, os, sys, math
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data', 'edsight')
ME, STATE = 'Andover School District', 'State of Connecticut'
CUR, PREV = '2025-26', '2024-25'
PEER_EXCLUDE = {'Norwich'}

NAVY, GOLD, INK, SOFT = '143862', '9A7400', '19222E', '4A566A'
H_FILL = PatternFill('solid', fgColor='E8EEF6')
ME_FILL = PatternFill('solid', fgColor='FBF2DC')
ST_FILL = PatternFill('solid', fgColor='E8EEF6')
THIN = Side(style='thin', color='D4DCE8')


def short(n):
    return n.replace(' School District', '').replace('Regional School District ', 'RSD ')


# ---------------------------------------------------------------- load
def load():
    lvl, hn, allc, grades = {}, {}, {}, {}
    ent, subj = {}, {}
    for r in csv.DictReader(open(os.path.join(DATA, 'performance-index', 'performance-index.csv'), encoding='utf-8')):
        k, y = (r['district'], r['year']), r['year']
        idx = float(r['index']) if r['index'] else None
        cnt = float(r['count']) if r['count'] else None
        if r['entityType']:
            ent[r['district']] = r['entityType']
        if r['studentGroup'] in ('District', 'State'):
            lvl.setdefault(k, {})[r['subject']] = (idx, cnt)
            subj.setdefault(k, {})[r['subject']] = (idx, cnt)
            if r['subject'] == 'ELA':
                allc[k] = cnt
        if r['category'] == 'High Needs (F/R, EL or SWD)' and r['studentGroup'] == 'High Needs' and r['subject'] == 'ELA':
            hn[k] = cnt
        if r['category'] == 'Grade' and r['subject'] == 'ELA' and cnt:
            grades.setdefault(k, set()).add(r['studentGroup'])
    ppe, ppef = {}, {}
    funcs = []
    for r in csv.DictReader(open(os.path.join(DATA, 'per-pupil-expenditures', 'per-pupil-expenditures.csv'), encoding='utf-8')):
        if not r['ppe']:
            continue
        k = (r['district'], r['year'])
        if r['function'] not in funcs:
            funcs.append(r['function'])
        ppef.setdefault(k, {})[r['function']] = float(r['ppe'])
        if r['function'] == 'Total':
            ppe[k] = float(r['ppe'])
    growth, acct = {}, {}
    for r in csv.DictReader(open(os.path.join(DATA, 'accountability', 'accountability.csv'), encoding='utf-8')):
        if r['level'] == 'school':
            continue
        k = (r['district'], r['year'])

        def f(c):
            try:
                return float(r[c])
            except Exception:
                return None
        e, m = f('Ind2ELA_All_Rate'), f('Ind2Math_All_Rate')
        growth[k] = (e + m) / 2 if (e is not None and m is not None) else None
        acct[k] = dict(index=f('outcomeRatePct'), academic=f('academicPct'), possible=f('possiblePoints'))
    peers = set()
    for r in csv.DictReader(open(os.path.join(ROOT, 'reports', 'aes', 'peer-spending',
                                              'peer_comparison_2024-25.csv'), encoding='utf-8-sig')):
        if r['District'] not in PEER_EXCLUDE:
            peers.add(r['District'] + ' School District')

    def weighted(s):
        v = [t for _, t in s.items() if t[0] is not None and t[1]]
        return sum(a * c for a, c in v) / sum(c for _, c in v) if len(v) >= 2 else None

    years = sorted({y for (_, y) in lvl})
    rows = {}
    for (d, y), s in lvl.items():
        if d != STATE and ent.get(d) not in ('local', 'regional'):
            continue
        g = grades.get((d, y), set())
        rows[(d, y)] = dict(
            district=d, name=short(d), year=y,
            level=weighted(s),
            eq=(sum(s[k][0] for k in ('ELA', 'Math', 'Science')) / 3
                if all(s.get(k, (None,))[0] is not None for k in ('ELA', 'Math', 'Science')) else None),
            ela=s.get('ELA', (None, None))[0], math=s.get('Math', (None, None))[0],
            sci=s.get('Science', (None, None))[0],
            nELA=s.get('ELA', (None, None))[1], nSci=s.get('Science', (None, None))[1],
            n=allc.get((d, y)),
            hn=(100 * hn[(d, y)] / allc[(d, y)]) if (hn.get((d, y)) and allc.get((d, y))) else None,
            hnCount=hn.get((d, y)),
            growth=growth.get((d, y)), ppe=ppe.get((d, y)), funcs=ppef.get((d, y), {}),
            acct=acct.get((d, y), {}),
            span=('K-12' if '11' in g else ('K-8' if ({'07', '08'} & g) else 'K-6')),
            peer=(d in peers), me=(d == ME), state=(d == STATE))
    # need-adjusted residual, against a fresh statewide fit each year
    for y in years:
        g = [r for r in rows.values() if r['year'] == y and not r['state']
             and r['level'] is not None and r['hn'] is not None]
        if len(g) < 30:
            continue
        n = len(g)
        mx = sum(r['hn'] for r in g) / n
        my = sum(r['level'] for r in g) / n
        sxx = sum((r['hn'] - mx) ** 2 for r in g)
        b = sum((r['hn'] - mx) * (r['level'] - my) for r in g) / sxx
        a = my - b * mx
        for r in g:
            r['predicted'] = a + b * r['hn']
            r['residual'] = r['level'] - r['predicted']
    return rows, years, funcs


# ---------------------------------------------------------------- sheet helpers
def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=NAVY, size=10)
        cell.fill = H_FILL
        cell.alignment = Alignment(vertical='bottom', wrap_text=True)
        cell.border = Border(bottom=THIN)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def widths(ws, spec):
    for i, w in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def table(wb, title, headers, rows, wid, fmts=None, highlight=True, autofilter=True):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r['cells'] if isinstance(r, dict) else r)
    style_header(ws, len(headers))
    widths(ws, wid)
    for i, r in enumerate(rows, start=2):
        tag = r.get('tag') if isinstance(r, dict) else None
        if highlight and tag:
            fill = ME_FILL if tag == 'me' else ST_FILL
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = fill
                ws.cell(row=i, column=c).font = Font(bold=True, size=10)
    if fmts:
        for col, fmt in fmts.items():
            for i in range(2, len(rows) + 2):
                ws.cell(row=i, column=col).number_format = fmt
    if autofilter and rows:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'
    return ws


N1, N2, PC, MONEY = '0.0', '0.00', '0.0"%"', '$#,##0'


# ---------------------------------------------------------------- main
def main():
    if len(sys.argv) < 2:
        sys.exit('usage: facts-workbook.py <target-folder>')
    out_dir = sys.argv[1] if os.path.isabs(sys.argv[1]) else os.path.join(ROOT, sys.argv[1])
    rows, years, funcs = load()
    districts = sorted({r['district'] for r in rows.values() if not r['state']})
    wb = Workbook()

    def get(d, y):
        return rows.get((d, y), {})

    def competition_rank(value, values):
        # Ties share the better rank (1 + how many are strictly higher). Published
        # indices carry one decimal, so ties are common, and a rank taken from list
        # position would depend on how the rows happened to be ordered.
        return 1 + sum(v > value + 1e-9 for v in values)

    def rank_of(d, y, key):
        g = [r for r in rows.values() if r['year'] == y and not r['state'] and r.get(key) is not None]
        me = next((r for r in g if r['district'] == d), None)
        return (competition_rank(me[key], [r[key] for r in g]) if me else None, len(g))

    # ---------------------------------------------------------------- About
    ws = wb.active
    ws.title = 'About'
    widths(ws, [3, 30, 104])
    A = []

    def line(label, text='', style=None):
        A.append((label, text, style))

    line('Andover and Connecticut school results', '', 'title')
    line('Companion data for the-facts edition "The Largest One-Year Gain in the State"', '', 'sub')
    line('')
    line('What this is', '', 'h')
    line('', 'Every figure cited in that edition, and the underlying district-by-district data it was drawn from, '
             'so that any number in it can be checked. Nothing here is original measurement: it is an extract and '
             'reshaping of public data published by the Connecticut State Department of Education.')
    line('', 'The "Key figures" sheet lists each number used in the edition and names the sheet it can be verified on.')
    line('')
    line('Where the data comes from', '', 'h')
    line('CT EdSight', 'public-edsight.ct.gov — the state Department of Education\'s public reporting portal. '
                       'Three of its reports are used here, each pulled in full rather than sampled.')
    line('1. Performance Index',
         'School, district and state performance indices by subject and student group, 0-100. '
         'Source assessments: Smarter Balanced (English and math, grades 3-8), SAT (grade 11), the NGSS science '
         'test (grades 5, 8, 11) and the Connecticut Alternate Assessments. RELEASED IN LATE AUGUST each year, '
         'so 2025-26 reflects assessments given in spring 2026. Years available: 2014-15 to 2025-26, with no '
         '2019-20 or 2020-21 (Connecticut held a federal accountability waiver for the pandemic years).')
    line('2. Next Generation Accountability',
         'The state accountability system: twelve indicators, of which indicator 2 is academic growth. '
         'RELEASED AROUND NOVEMBER following the school year, so the most recent available is 2024-25 and '
         '2025-26 growth is not yet published. Years available: 2014-15 to 2024-25, same pandemic gap.')
    line('3. Per Pupil Expenditures by Function (District)',
         'District spending per pupil, split across eleven functions plus a total. Reported by districts through '
         'the Education Financial System; the state notes the data is NOT FULLY AUDITED and audits may change it. '
         'Most recent available is 2024-25. Years available: 2017-18 to 2024-25.')
    line('')
    line('What the measures mean', '', 'h')
    line('Performance index',
         'The state\'s 0-100 summary of how students performed on the assessments. EDSIGHT PUBLISHES IT SEPARATELY '
         'FOR EACH SUBJECT — English language arts, mathematics and science — and does not publish a single '
         'combined figure. Connecticut\'s stated target is 75 in each subject. The ELA, math and science figures in '
         'this workbook are exactly as published.')
    line('Combined index (ours)',
         'Where a sheet shows one figure across all three subjects, it is THIS WORKBOOK\'S OWN CALCULATION, not an '
         'EdSight number: the three subject indices weighted by the number of students who sat each, so science — '
         'tested in only three grades — is not given a third of the weight. The choice matters: from 2024-25 to '
         '2025-26, Andover\'s combined gain ranks first weighted by students, but sixth weighted equally, because '
         'science fell while English and math rose. Both versions are on the "Change" sheet.')
    line('Academic growth',
         'Indicator 2 of the accountability system, and a longitudinal measure: Smarter Balanced is vertically '
         'scaled, so each student in grades 4-8 is given an individual growth target from their OWN prior-year '
         'score. The district figure is the average percentage of that target achieved, capped at 110% per student. '
         'It is NOT the share of students who met their target — the state reports that separately as the "growth '
         'rate" and does not use it for accountability. English and math only, grades 4-8 only, minimum 20 matched '
         'students.')
    line('High needs',
         'A student eligible for free or reduced-price meals, or an English learner, or a student with a '
         'disability. The share shown is high-needs students as a percentage of students tested. The Connecticut '
         'figure is the statewide student share, which sits well above most towns because high-need students are '
         'concentrated in the cities.')
    line('Need-adjusted residual',
         'How far a district sits above or below what its high-needs share alone predicts, from a least-squares '
         'fit recomputed across all town districts SEPARATELY IN EACH YEAR. Zero means exactly as predicted. '
         'Because the fit is redone each year, statewide movements — including pandemic learning loss — are '
         'already netted out, so the residual measures position relative to comparable districts, not absolute '
         'performance.')
    line('Per-pupil spending',
         'Total per-pupil expenditure across all functions. Excludes debt, capital beyond equipment, adult '
         'education, community services, non-local food service and the state\'s teacher retirement contributions, '
         'so it will NOT match a town\'s appropriated school budget. Note that the denominator varies by function: '
         'student transportation is per transported pupil, not per enrolled pupil.')
    line('')
    line('Who is included', '', 'h')
    line('165 town districts',
         'All comparisons cover Connecticut\'s local and regional town school districts — 165 in 2024-25, 164 in '
         '2025-26. Charter districts, the six regional educational service centers, the technical high school '
         'system and state-agency schools are excluded throughout: they serve different populations under '
         'different rules, and the service centers in particular top every per-pupil spending ranking, which '
         'would distort any cross-town comparison.')
    line('45-town peer group',
         'The comparison set from this site\'s earlier report on peer spending: districts that run only '
         'elementary grades and send older students elsewhere for high school. Norwich meets that structural test '
         'but was excluded from the analysis in that report, as a city district several times the size of any '
         'other member, and is excluded here on the same basis.')
    line('')
    line('Cautions worth carrying', '', 'h')
    line('Small numbers',
         'Andover tests about 109 students. One child is worth roughly a full percentage point of any subgroup '
         'figure, and among districts testing fewer than 200 the typical year-over-year swing is about two and a '
         'half times that of districts testing 600 or more. Single-year movements in small districts should be '
         'read with that in mind, in either direction.')
    line('Suppression',
         'The state withholds figures based on too few students. Those appear as blanks here, not zeros. Science '
         'in particular is suppressed for many small districts.')
    line('Grade span',
         'Districts that do not test older grades score structurally higher, because performance falls with grade '
         'level: statewide, the English index runs about 68 in grades 3-6 and about 61 in grades 7-11. Andover is '
         'a K-6 district. Comparisons against K-12 districts are not like-for-like on this measure.')
    line('What this data cannot settle',
         'Per-pupil spending explains very little of the difference between districts, and once student need and '
         'grade span are accounted for, no statistically detectable amount. That is not evidence that spending '
         'does not matter — it is evidence that what varies between Connecticut towns is mostly district size and '
         'geography rather than educational effort, and that these figures cannot resolve a budget argument in '
         'either direction.')
    line('')
    line('Provenance', '', 'h')
    line('Compiled by', 'Scott Sauyet, Andover, Connecticut — andoverct.info')
    line('Method', 'Data pulled directly from EdSight\'s published exports and reshaped without adjustment. '
                   'The extraction and this workbook are generated by scripts kept with the site source, so the '
                   'figures here, the charts in the edition and the text all derive from one extract.')
    line('Verify anything', 'Every figure can be reproduced from EdSight itself at public-edsight.ct.gov.')

    r = 1
    for label, text, style in A:
        if style == 'title':
            ws.cell(row=r, column=2, value=label).font = Font(bold=True, size=17, color=NAVY)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            r += 1
            continue
        if style == 'sub':
            ws.cell(row=r, column=2, value=label).font = Font(italic=True, size=11, color=SOFT)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            r += 2
            continue
        if style == 'h':
            ws.cell(row=r, column=2, value=label).font = Font(bold=True, size=12, color=GOLD)
            r += 1
            continue
        if label:
            c = ws.cell(row=r, column=2, value=label)
            c.font = Font(bold=True, size=10, color=INK)
            c.alignment = Alignment(vertical='top', wrap_text=True)
        if text:
            c = ws.cell(row=r, column=3, value=text)
            c.font = Font(size=10, color=INK)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            ws.row_dimensions[r].height = max(15, 12.2 * (len(text) // 92 + 1))
        r += 1
    ws.sheet_view.showGridLines = False

    # ---------------------------------------------------------------- Key figures
    a_cur, a_prev = get(ME, CUR), get(ME, PREV)
    s_cur, s_prev = get(STATE, CUR), get(STATE, PREV)

    def median(v):
        v = sorted(v)
        return (v[(len(v) - 1) // 2] + v[len(v) // 2]) / 2

    def change_stats(k):
        deltas = []
        for d in districts:
            x, y = get(d, PREV), get(d, CUR)
            if x.get(k) is not None and y.get(k) is not None:
                deltas.append((d, y[k] - x[k]))
        deltas.sort(key=lambda t: -t[1])
        vals = [v for _, v in deltas]
        names = [d for d, _ in deltas]
        return dict(deltas=deltas, n=len(vals), median=median(vals),
                    rank=competition_rank(dict(deltas)[ME], vals),
                    change=dict(deltas)[ME], next=deltas[1][1] if names[0] == ME else deltas[0][1],
                    three=sum(v >= 3 for v in vals), improved=100 * sum(v > 0 for v in vals) / len(vals))

    stats = {k: change_stats(k) for k in ('ela', 'math', 'level', 'eq')}

    def prior_best_gain(k):
        hist = [(y, get(ME, y)[k]) for y in years if get(ME, y).get(k) is not None]
        gains = [(hist[i][1] - hist[i - 1][1], hist[i - 1][0], hist[i][0]) for i in range(1, len(hist) - 1)]
        return max(gains)

    def corr(y):
        g = [r for r in rows.values() if r['year'] == y and not r['state']
             and r['level'] is not None and r['hn'] is not None]
        n = len(g)
        mx = sum(r['hn'] for r in g) / n
        my = sum(r['level'] for r in g) / n
        sxx = sum((r['hn'] - mx) ** 2 for r in g)
        syy = sum((r['level'] - my) ** 2 for r in g)
        sxy = sum((r['hn'] - mx) * (r['level'] - my) for r in g)
        return sxy / math.sqrt(sxx * syy), n

    def peer_rank(y):
        g = sorted([r for r in rows.values() if r['year'] == y and r['peer'] and r['level'] is not None],
                   key=lambda r: -r['level'])
        return [r['district'] for r in g].index(ME) + 1, len(g)

    def swing_sd(lo, hi):
        v = [dv for d, dv in stats['level']['deltas'] if lo <= (get(d, CUR).get('n') or 0) < hi]
        m = sum(v) / len(v)
        return math.sqrt(sum((x - m) ** 2 for x in v) / (len(v) - 1))

    def ranks(k):
        a, b = rank_of(ME, PREV, k), rank_of(ME, CUR, k)
        return f'{a[0]} of {a[1]} to {b[0]} of {b[1]}'

    r25, n25 = corr(PREV)
    r26, n26 = corr(CUR)
    peers_cur = [d for d in districts if get(d, CUR).get('peer')]
    S_ELA, S_M, S_C, S_E = stats['ela'], stats['math'], stats['level'], stats['eq']
    pg_ela, pg_m = prior_best_gain('ela'), prior_best_gain('math')
    pr_prev, pr_cur = peer_rank(PREV), peer_rank(CUR)
    CH = 'Change 24-25 to 25-26'
    AC = 'Andover and Connecticut'
    DY = 'Districts 2024-25 / 2025-26'
    a18 = get(ME, '2018-19')
    kf = [
        ('PUBLISHED BY EDSIGHT, BY SUBJECT', '', ''),
        ('Andover ELA index, 2024-25 to 2025-26', f"{a_prev['ela']:.1f} to {a_cur['ela']:.1f}", CH),
        ('Andover ELA one-year change', f"{S_ELA['change']:+.1f}", CH),
        ('Rank of that change among town districts', f"{S_ELA['rank']} of {S_ELA['n']}", CH),
        ('Next largest ELA gain', f"{S_ELA['next']:+.1f}", CH),
        ('Median ELA change across town districts', f"{S_ELA['median']:+.2f}", CH),
        ('Districts gaining 3+ points in ELA', S_ELA['three'], CH),
        ('Share of districts improving in ELA', f"{S_ELA['improved']:.0f}%", CH),
        ('Andover math index, 2024-25 to 2025-26', f"{a_prev['math']:.1f} to {a_cur['math']:.1f}", CH),
        ('Andover math one-year change', f"{S_M['change']:+.1f}", CH),
        ('Rank of that change among town districts', f"{S_M['rank']} of {S_M['n']}", CH),
        ('Next largest math gain', f"{S_M['next']:+.1f}", CH),
        ('Median math change across town districts', f"{S_M['median']:+.2f}", CH),
        ('Districts gaining 3+ points in math', S_M['three'], CH),
        ('Share of districts improving in math', f"{S_M['improved']:.0f}%", CH),
        ('Andover science index, 2024-25 to 2025-26', f"{a_prev['sci']:.1f} to {a_cur['sci']:.1f}", CH),
        ('Connecticut ELA index, 2024-25 to 2025-26', f"{s_prev['ela']:.1f} to {s_cur['ela']:.1f}", AC),
        ('Connecticut math index, 2024-25 to 2025-26', f"{s_prev['math']:.1f} to {s_cur['math']:.1f}", AC),
        ('State target, each subject', 75, AC),
        ('Andover ELA rank among town districts', ranks('ela'), DY),
        ('Andover math rank among town districts', ranks('math'), DY),
        ('Andover largest prior one-year ELA gain', f"{pg_ela[0]:+.1f} ({pg_ela[1]} to {pg_ela[2]})", 'ELA by year'),
        ('Andover largest prior one-year math gain', f"{pg_m[0]:+.1f} ({pg_m[1]} to {pg_m[2]})", 'Math by year'),
        ('Andover ELA and math, 2018-19', f"{a18['ela']:.1f} and {a18['math']:.1f}", AC),
        ('COMBINED ACROSS SUBJECTS: OUR CALCULATION, NOT AN EDSIGHT FIGURE', '', ''),
        ('Andover combined index (weighted by students tested)', f"{a_prev['level']:.1f} to {a_cur['level']:.1f}", CH),
        ('Andover combined change, and its rank', f"{S_C['change']:+.1f}, {S_C['rank']} of {S_C['n']}", CH),
        ('Same, subjects weighted equally, and its rank', f"{S_E['change']:+.1f}, {S_E['rank']} of {S_E['n']}", CH),
        ('Connecticut combined index', f"{s_prev['level']:.1f} to {s_cur['level']:.1f}", AC),
        ('Andover high-needs share', f"{a_prev['hn']:.1f}% to {a_cur['hn']:.1f}%", AC),
        ('Andover need-adjusted residual', f"{a_prev['residual']:+.1f} to {a_cur['residual']:+.1f}", AC),
        ('Gain surviving the need adjustment', f"{a_cur['residual'] - a_prev['residual']:+.1f}", AC),
        ('Andover combined rank among town districts', ranks('level'), DY),
        ('Andover combined rank within the peer group', f"{pr_prev[0]} of {pr_prev[1]} to {pr_cur[0]} of {pr_cur[1]}", 'Peer group'),
        ('Combined index vs high-needs share, 2024-25', f'r = {r25:+.3f} (r² {r25 * r25:.2f}, n = {n25})', 'Districts 2024-25'),
        ('Combined index vs high-needs share, 2025-26', f'r = {r26:+.3f} (r² {r26 * r26:.2f}, n = {n26})', 'Districts 2025-26'),
        ('Year-to-year swing in the combined index (s.d.)',
         f"{swing_sd(0, 200):.2f} under 200 tested vs {swing_sd(600, 10 ** 9):.2f} at 600+", CH),
        ('OTHER MEASURES', '', ''),
        ('Andover academic growth, 2024-25', round(a_prev['growth'], 1), 'Districts 2024-25'),
        ('Connecticut academic growth, 2024-25', round(s_prev['growth'], 1), AC),
        ('Andover per-pupil spending, 2024-25', round(a_prev['ppe']), 'Spending by function 2024-25'),
        ('Connecticut per-pupil spending, 2024-25', round(s_prev['ppe']), 'Spending by function 2024-25'),
        ('Andover students tested, 2025-26', int(a_cur['n']), 'Districts 2025-26'),
        ('Peer group size, 2025-26', len(peers_cur), 'Peer group'),
    ]
    ws = wb.create_sheet('Key figures')
    ws.append(['Figure cited in the edition', 'Value', 'Check it on this sheet'])
    for k, v, sheet in kf:
        ws.append([k, v, sheet])
    style_header(ws, 3)
    widths(ws, [52, 26, 30])
    # The value column mixes numbers and text; Excel would right-align the
    # numbers and left-align the text, so pin it left to read as one column.
    for (cell,) in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        cell.alignment = Alignment(horizontal='left')
    ws.sheet_view.showGridLines = False

    # ---------------------------------------------------------------- Andover and Connecticut
    hdr = ['Year', 'Andover combined index (ours)', 'Andover ELA', 'Andover math', 'Andover science',
           'Andover growth', 'Andover high needs %', 'Andover tested', 'Andover per-pupil $',
           'Predicted from need', 'Residual', 'CT combined index (ours)', 'CT growth', 'CT high needs %',
           'CT per-pupil $', 'CT ELA', 'CT math', 'CT science']
    body = []
    for y in years:
        a, s = get(ME, y), get(STATE, y)
        body.append([y, a.get('level'), a.get('ela'), a.get('math'), a.get('sci'), a.get('growth'),
                     a.get('hn'), a.get('n'), a.get('ppe'), a.get('predicted'), a.get('residual'),
                     s.get('level'), s.get('growth'), s.get('hn'), s.get('ppe'),
                     s.get('ela'), s.get('math'), s.get('sci')])
    ws = table(wb, 'Andover and Connecticut', hdr, body, [10] + [13] * 17,
               fmts={i: N1 for i in list(range(2, 8)) + [10, 11, 12, 13, 14, 16, 17, 18]}, autofilter=False)
    for i in range(2, len(body) + 2):
        ws.cell(row=i, column=9).number_format = MONEY
        ws.cell(row=i, column=15).number_format = MONEY
        ws.cell(row=i, column=7).number_format = PC
        ws.cell(row=i, column=14).number_format = PC

    # ---------------------------------------------------------------- per-year district sheets
    for y, label in ((CUR, 'Districts 2025-26'), (PREV, 'Districts 2024-25')):
        hdr = ['District', 'Type', 'Grade span', 'Peer group', 'ELA', 'Math', 'Science',
               'Combined index (ours)', 'Academic growth', 'High needs %', 'High-needs students',
               'Students tested', 'Per-pupil spending', 'Predicted from need', 'Residual',
               'Rank (ELA)', 'Rank (math)', 'Rank (combined)']
        g = [r for r in rows.values() if r['year'] == y and not r['state']]
        g.sort(key=lambda r: -(r['level'] if r['level'] is not None else -1e9))
        body = []
        for r in g:
            d = r['district']
            body.append(dict(tag='me' if r['me'] else None, cells=[
                r['name'], 'Regional' if r['name'].startswith('RSD') else 'Local', r['span'],
                'yes' if r['peer'] else '', r['ela'], r['math'], r['sci'], r['level'], r['growth'],
                r['hn'], r['hnCount'], r['n'], r['ppe'], r.get('predicted'), r.get('residual'),
                rank_of(d, y, 'ela')[0], rank_of(d, y, 'math')[0], rank_of(d, y, 'level')[0]]))
        st = get(STATE, y)
        body.insert(0, dict(tag='state', cells=[
            'State of Connecticut', 'Statewide', '', '', st.get('ela'), st.get('math'), st.get('sci'),
            st.get('level'), st.get('growth'), st.get('hn'), st.get('hnCount'), st.get('n'), st.get('ppe'),
            '', '', '', '', '']))
        ws = table(wb, label, hdr, body, [30, 10, 11, 11, 9, 9, 9, 15, 14, 12, 15, 12, 15, 15, 11, 10, 11, 14],
                   fmts={i: N1 for i in (5, 6, 7, 8, 9, 14, 15)})
        for i in range(2, len(body) + 2):
            ws.cell(row=i, column=10).number_format = PC
            ws.cell(row=i, column=13).number_format = MONEY

    # ---------------------------------------------------------------- change
    # Sorted by the published ELA change; every rank column is its own ranking,
    # so the sheet can be re-sorted on any of them.
    hdr = ['District', 'Peer group',
           'ELA 2024-25', 'ELA 2025-26', 'ELA change', 'ELA rank',
           'Math 2024-25', 'Math 2025-26', 'Math change', 'Math rank',
           'Science change',
           'Combined change, weighted by students (ours)', 'Rank',
           'Combined change, subjects weighted equally', 'Rank ',
           'Change in high needs (pts)', 'Change in need-adjusted residual']
    rk = {k: {d: competition_rank(v, [x for _, x in stats[k]['deltas']]) for d, v in stats[k]['deltas']}
          for k in stats}
    body = []
    for d, _ in stats['ela']['deltas']:
        x, y2 = get(d, PREV), get(d, CUR)

        def ch(k):
            return (y2[k] - x[k]) if (x.get(k) is not None and y2.get(k) is not None) else None
        body.append(dict(tag='me' if d == ME else None, cells=[
            short(d), 'yes' if y2.get('peer') else '',
            x.get('ela'), y2.get('ela'), ch('ela'), rk['ela'].get(d),
            x.get('math'), y2.get('math'), ch('math'), rk['math'].get(d),
            ch('sci'), ch('level'), rk['level'].get(d), ch('eq'), rk['eq'].get(d),
            ch('hn'), ch('residual')]))
    table(wb, 'Change 24-25 to 25-26', hdr, body,
          [30, 11, 11, 11, 10, 9, 12, 12, 11, 10, 13, 20, 7, 20, 7, 16, 18],
          fmts={i: N1 for i in (3, 4, 5, 7, 8, 9, 11, 12, 14, 16, 17)})

    # ---------------------------------------------------------------- trends
    for key, label, fmt in (('ela', 'ELA by year', N1), ('math', 'Math by year', N1),
                            ('level', 'Combined index by year', N1), ('growth', 'Growth by year', N1),
                            ('hn', 'High needs by year', PC), ('residual', 'Residual by year', N1)):
        ys = [y for y in years if any(rows.get((d, y), {}).get(key) is not None for d in districts)]
        hdr = ['District', 'Peer group'] + ys
        body = []
        for d in ['__state__'] + districts:
            if d == '__state__':
                if key == 'residual':
                    continue
                body.append(dict(tag='state', cells=['State of Connecticut', ''] +
                                 [get(STATE, y).get(key) for y in ys]))
                continue
            body.append(dict(tag='me' if d == ME else None,
                             cells=[short(d), 'yes' if get(d, CUR).get('peer') or get(d, PREV).get('peer') else ''] +
                                   [get(d, y).get(key) for y in ys]))
        table(wb, label, hdr, body, [30, 11] + [10] * len(ys),
              fmts={i: fmt for i in range(3, 3 + len(ys))})

    # ---------------------------------------------------------------- spending by function
    order = [f for f in funcs if f != 'Total'] + ['Total']
    hdr = ['District', 'Peer group'] + order
    body = []
    st = get(STATE, PREV)
    body.append(dict(tag='state', cells=['State of Connecticut', ''] + [st.get('funcs', {}).get(f) for f in order]))
    for d in districts:
        r = get(d, PREV)
        body.append(dict(tag='me' if d == ME else None,
                         cells=[short(d), 'yes' if r.get('peer') else ''] + [r.get('funcs', {}).get(f) for f in order]))
    table(wb, 'Spending by function 2024-25', hdr, body, [30, 11] + [16] * len(order),
          fmts={i: MONEY for i in range(3, 3 + len(order))})

    # ---------------------------------------------------------------- peer group
    hdr = ['District', 'Grade span', 'ELA 2025-26', 'ELA change', 'Math 2025-26', 'Math change',
           'Combined index 2025-26 (ours)', 'Combined change (ours)',
           '2024-25 growth', 'High needs %', 'Students tested', 'Per-pupil spending 2024-25']
    body = []

    def chg(c, p, k):
        return (c[k] - p[k]) if (c.get(k) is not None and p.get(k) is not None) else None
    for d in sorted(peers_cur, key=lambda d: -(get(d, CUR).get('level') or -1e9)):
        c, p = get(d, CUR), get(d, PREV)
        body.append(dict(tag='me' if d == ME else None, cells=[
            short(d), c.get('span'), c.get('ela'), chg(c, p, 'ela'), c.get('math'), chg(c, p, 'math'),
            c.get('level'), chg(c, p, 'level'), p.get('growth'), c.get('hn'), c.get('n'), p.get('ppe')]))
    ws = table(wb, 'Peer group', hdr, body, [30, 12, 12, 12, 13, 12, 16, 15, 14, 13, 14, 16],
               fmts={i: N1 for i in (3, 4, 5, 6, 7, 8, 9)})
    for i in range(2, len(body) + 2):
        ws.cell(row=i, column=10).number_format = PC
        ws.cell(row=i, column=12).number_format = MONEY

    out = os.path.join(out_dir, 'andover-edsight-data.xlsx')
    wb.save(out)
    size = os.path.getsize(out) / 1024
    print(f'wrote {out}  ({size:.0f} KB)')
    print('sheets: ' + ', '.join(wb.sheetnames))


if __name__ == '__main__':
    main()
